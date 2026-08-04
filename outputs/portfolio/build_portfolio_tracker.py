#!/usr/bin/env python3
"""Live portfolio tracker: Dashboard, Prices, Estimates, Targets, a tab per holding."""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "/tmp/claude-0/-home-user/de029611-12c7-500a-9a8b-f4dd5e8389f7/scratchpad/build/Portfolio_Tracker.xlsx"

# ============================================================== source data
# Cost basis / shares: Trading 212 Never Sell pie export, 4 Aug 2026.
# Prices, 52w range: FMP EOD close 3 Aug 2026 (last close before the export).
# EPS + targets: FMP analyst consensus, pulled 4 Aug 2026.
GBPUSD = 1.34499          # FMP forex quote GBPUSD, 4 Aug 2026
PRICE_DATE = "2026-08-03"
Y1, Y2, Y3 = "FY2027", "FY2028", "FY2029"

HOLDINGS = [
    dict(tkr="AMZN", name="Amazon.com, Inc.", sleeve="Never Sell",
         invested=12204.39, snapshot_value=13618.02, shares=66.0103,
         price=284.02, prev=271.58, hi52=284.02, lo52=198.79,
         fy_end="December", last_fy="FY2025", eps_last=7.15268,
         eps=[11.70249, 10.42302, 13.57301, 16.20478, 19.82384],   # FY26..FY30
         tgt=322.68, tgt_med=325.0, tgt_lo=175.0, tgt_hi=390.0, n_analysts=41,
         dps=0.0),
    dict(tkr="MSFT", name="Microsoft Corporation", sleeve="Never Sell",
         invested=5800.47, snapshot_value=7250.23, shares=20.4311,
         price=487.65, prev=464.72, hi52=542.07, lo52=352.83,
         fy_end="June", last_fy="FY2026", eps_last=17.00426,
         eps=[17.00426, 19.63054, 23.25411, 28.17438, 35.4625],
         tgt=541.88, tgt_med=531.5, tgt_lo=400.0, tgt_hi=680.0, n_analysts=29,
         dps=3.32),
    dict(tkr="UBER", name="Uber Technologies, Inc.", sleeve="Never Sell",
         invested=5608.26, snapshot_value=5831.61, shares=109.6538,
         price=71.61, prev=70.36, hi52=100.10, lo52=65.94,
         fy_end="December", last_fy="FY2025", eps_last=5.35736,
         eps=[3.34187, 4.47576, 5.50281, 6.51504, 6.96167],
         tgt=102.25, tgt_med=105.0, tgt_lo=72.0, tgt_hi=125.0, n_analysts=30,
         dps=0.0),
]

GOAL = 1000000  # £1m target

# ================================================================= styling
FONT = "Arial"
BLUE  = Font(name=FONT, size=10, color="0000FF")
BLACK = Font(name=FONT, size=10)
GREEN = Font(name=FONT, size=10, color="008000")
LABEL = Font(name=FONT, size=10)
BOLD  = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="FFFFFF")
SECT  = Font(name=FONT, size=11, bold=True, color="FFFFFF")
HEADF = Font(name=FONT, size=10, bold=True, color="FFFFFF")
NOTE  = Font(name=FONT, size=9, italic=True, color="808080")
BIG   = Font(name=FONT, size=12, bold=True)

NAVY   = PatternFill("solid", fgColor="1F3864")
SLATE  = PatternFill("solid", fgColor="44546A")
YELLOW = PatternFill("solid", fgColor="FFFF00")
LGREY  = PatternFill("solid", fgColor="F2F2F2")
BANNER = PatternFill("solid", fgColor="D9E1F2")
GOLD   = PatternFill("solid", fgColor="FFF2CC")

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

GBP  = '£#,##0.00;(£#,##0.00);"-"'
GBP0 = '£#,##0;(£#,##0);"-"'
USD  = '$#,##0.00;($#,##0.00);"-"'
PCT  = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
SHR  = '#,##0.0000;(#,##0.0000);"-"'
MULT = '0.0"x"'
NUM2 = '0.00'
INT  = '#,##0'

wb = openpyxl.Workbook()


def sc(ws, coord, value=None, font=BLACK, fill=None, fmt=None, align=None,
       border=False, comment=None, wrap=False):
    c = ws[coord]
    if value is not None:
        c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align or wrap:
        c.alignment = Alignment(horizontal=align or "general",
                                vertical="center", wrap_text=wrap)
    if border:
        c.border = BOX
    if comment:
        c.comment = Comment(comment, "Portfolio Tracker")
    return c


def banner(ws, row, text, last_col="F", fill=NAVY, font=TITLE, height=22):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    sc(ws, f"A{row}", text, font=font, fill=fill, align="left")
    for col in range(1, openpyxl.utils.column_index_from_string(last_col) + 1):
        ws.cell(row=row, column=col).fill = fill
    ws.row_dimensions[row].height = height


def section(ws, row, text, last_col="F"):
    banner(ws, row, text, last_col, fill=SLATE, font=SECT, height=18)


def headers(ws, row, cols, widths=None):
    for i, h in enumerate(cols):
        col = get_column_letter(i + 1)
        sc(ws, f"{col}{row}", h, font=HEADF, fill=SLATE, border=True,
           align="center", wrap=True)
    ws.row_dimensions[row].height = 30


# ==================================================== 1. Estimates (data hub)
es = wb.create_sheet("Estimates")
es.sheet_properties.tabColor = "7030A0"
es.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMNOPQ",
                  [10, 24, 11, 12, 13, 12, 12, 12, 12, 12, 12, 13, 12, 11, 11, 13, 11]):
    es.column_dimensions[col].width = w

banner(es, 1, "CONSENSUS ESTIMATES  —  the EPS and price targets everything else is built on",
       last_col="Q")
sc(es, "A2", f"Source: FMP analyst consensus, pulled 4 August 2026. "
             f"Year 1 = {Y1}, Year 2 = {Y2}, Year 3 = {Y3}.", font=NOTE)

EHEAD = 4
headers(es, EHEAD, ["Ticker", "Company", "FY ends", "Last rep. FY", "EPS last rep.",
                    "FY2026E", f"{Y1}E (Y1)", f"{Y2}E (Y2)", f"{Y3}E (Y3)", "FY2030E",
                    "EPS CAGR Y1>Y3", "Target cons.", "Target med.", "Target low",
                    "Target high", "Implied P/E on Y1", "Analysts"])
es.freeze_panes = f"A{EHEAD+1}"

EROW = {}
r = EHEAD + 1
for h in HOLDINGS:
    EROW[h["tkr"]] = r
    sc(es, f"A{r}", h["tkr"], font=BOLD, border=True, align="center")
    sc(es, f"B{r}", h["name"], font=LABEL, border=True)
    sc(es, f"C{r}", h["fy_end"], font=BLUE, border=True, align="center")
    sc(es, f"D{r}", h["last_fy"], font=BLUE, border=True, align="center")
    sc(es, f"E{r}", h["eps_last"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment="Consensus EPS for the last completed fiscal year (USD, diluted).\n"
               "Source: FMP analyst consensus, 4 Aug 2026.")
    for i, col in enumerate("FGHIJ"):
        sc(es, f"{col}{r}", h["eps"][i], font=BLUE, fill=YELLOW, fmt=USD, border=True)
    sc(es, f"K{r}", f"=IFERROR((I{r}/G{r})^(1/2)-1,0)", font=BLACK, fmt=PCT, border=True)
    sc(es, f"L{r}", h["tgt"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment="Consensus analyst price target (USD).\nSource: FMP price-target-consensus, 4 Aug 2026.")
    sc(es, f"M{r}", h["tgt_med"], font=BLUE, fmt=USD, border=True)
    sc(es, f"N{r}", h["tgt_lo"], font=BLUE, fmt=USD, border=True)
    sc(es, f"O{r}", h["tgt_hi"], font=BLUE, fmt=USD, border=True)
    sc(es, f"P{r}", f"=IFERROR(L{r}/G{r},0)", font=BLACK, fmt=MULT, border=True,
       comment="What the street's target implies you should pay for Year 1 earnings.\n"
               "This is the starting value for your own target P/E on each holding tab.")
    sc(es, f"Q{r}", h["n_analysts"], font=BLUE, fmt=INT, border=True, align="center")
    r += 1
ELAST = r - 1

r += 1
sc(es, f"A{r}", "Refresh these once a quarter, after earnings. Everything downstream moves with them.", font=NOTE)
r += 1
sc(es, f"A{r}", "MSFT's fiscal year ends in June, so FY2026 is already reported; AMZN and UBER run to December.", font=NOTE)
r += 1
sc(es, f"A{r}", "AMZN FY2026E sits above FY2027E because the 2026 consensus carries a large one-off gain.", font=NOTE)
r += 1
sc(es, f"A{r}", "UBER FY2025 EPS is flattered by a one-off tax benefit; FY2026E is the cleaner run-rate.", font=NOTE)


# ================================================================ 2. Prices
pr = wb.create_sheet("Prices")
pr.sheet_properties.tabColor = "00B050"
pr.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKL",
                  [11, 24, 8, 13, 13, 12, 11, 13, 13, 12, 13, 30]):
    pr.column_dimensions[col].width = w

banner(pr, 1, "PRICE ENGINE  —  update the yellow cells and the whole workbook moves",
       last_col="L")
FX_ROW, DATE_ROW = 3, 4
sc(pr, "A3", "GBP/USD rate", font=BOLD)
sc(pr, "B3", GBPUSD, font=BLUE, fill=YELLOW, fmt='0.00000', align="center",
   comment="Source: FMP forex quote GBPUSD, 4 Aug 2026.\n"
           "Trading 212 converts at its own rate plus a 0.15% FX fee, so their GBP "
           "values can differ slightly from this workbook's.")
sc(pr, "C3", "<- every USD figure is converted at this rate", font=NOTE)
sc(pr, "A4", "Prices as at", font=BOLD)
sc(pr, "B4", PRICE_DATE, font=BLUE, fill=YELLOW, align="center")
sc(pr, "C4", "<- stamp the date/time you last refreshed", font=NOTE)

PHEAD = 6
headers(pr, PHEAD, ["Ticker", "Company", "Ccy", "Live price", "Prev close",
                    "Day chg", "Day chg %", "52w high", "52w low", "% off high",
                    "Price (GBP)", "Source"])
pr.freeze_panes = f"A{PHEAD+1}"

PROW = {}
r = PHEAD + 1
for h in HOLDINGS:
    PROW[h["tkr"]] = r
    sc(pr, f"A{r}", h["tkr"], font=BOLD, border=True, align="center")
    sc(pr, f"B{r}", h["name"], font=LABEL, border=True)
    sc(pr, f"C{r}", "USD", font=BLUE, border=True, align="center")
    sc(pr, f"D{r}", h["price"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment=f"Closing price {PRICE_DATE} (USD).\nSource: FMP end-of-day close.")
    sc(pr, f"E{r}", h["prev"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment="Prior session close, drives the day-change columns.")
    sc(pr, f"F{r}", f'=IF(E{r}="","",D{r}-E{r})', font=BLACK, fmt=USD, border=True)
    sc(pr, f"G{r}", f'=IFERROR(F{r}/E{r},"")', font=BLACK, fmt=PCT2, border=True)
    sc(pr, f"H{r}", h["hi52"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment="Highest closing price in the last 52 weeks (4 Aug 2025 - 3 Aug 2026).\n"
               "Closing basis, not intraday. Drives the trailing stop.")
    sc(pr, f"I{r}", h["lo52"], font=BLUE, fill=YELLOW, fmt=USD, border=True,
       comment="Lowest closing price in the last 52 weeks. Closing basis, not intraday.")
    sc(pr, f"J{r}", f'=IFERROR(D{r}/H{r}-1,"")', font=BLACK, fmt=PCT, border=True)
    sc(pr, f"K{r}", f'=IFERROR(IF(C{r}="USD",D{r}/$B${FX_ROW},D{r}),"")',
       font=BLACK, fmt=GBP, border=True)
    sc(pr, f"L{r}", f"FMP EOD close {PRICE_DATE}", font=NOTE, border=True)
    r += 1
PLAST = r - 1

red_f = Font(name=FONT, size=10, color="C00000")
green_f = Font(name=FONT, size=10, color="008000")
pr.conditional_formatting.add(f"F{PHEAD+1}:G{PLAST}",
    CellIsRule(operator="lessThan", formula=["0"], font=red_f))
pr.conditional_formatting.add(f"F{PHEAD+1}:G{PLAST}",
    CellIsRule(operator="greaterThan", formula=["0"], font=green_f))

# reconciliation against the pie export
r = PLAST + 2
section(pr, r, "RECONCILIATION  —  why this differs from your Trading 212 pie export", last_col="L")
r += 1
RHEAD = r
headers(pr, RHEAD, ["Ticker", "Pie export value", "This workbook", "Difference",
                    "Diff %", "Why"])
r += 1
RFIRST = r
for h in HOLDINGS:
    tkr = h["tkr"]
    sc(pr, f"A{r}", tkr, font=BOLD, border=True, align="center")
    sc(pr, f"B{r}", h["snapshot_value"], font=BLUE, fmt=GBP, border=True)
    sc(pr, f"C{r}", f"='{tkr}'!$B$14", font=GREEN, fmt=GBP, border=True)
    sc(pr, f"D{r}", f"=C{r}-B{r}", font=BLACK, fmt=GBP, border=True)
    sc(pr, f"E{r}", f"=IFERROR(D{r}/B{r},0)", font=BLACK, fmt=PCT, border=True)
    r += 1
RLAST = r - 1
sc(pr, f"A{r}", "TOTAL", font=BOLD, fill=BANNER, border=True)
for col in "BCD":
    sc(pr, f"{col}{r}", f"=SUM({col}{RFIRST}:{col}{RLAST})", font=BOLD, fill=BANNER,
       fmt=GBP, border=True)
sc(pr, f"E{r}", f"=IFERROR(D{r}/B{r},0)", font=BOLD, fill=BANNER, fmt=PCT, border=True)
sc(pr, f"F{r}", None, fill=BANNER, border=True)
r += 2
for txt in [
    "The pie export was taken 4 Aug at 11:15 UTC but its prices lag the 3 Aug US close.",
    "AMZN and MSFT both jumped hard on 30-31 July results, so the export understates them by about 2%.",
    "UBER matches to 0.1%, which is what tells you the gap is stale prices and not a bad FX rate.",
    "Your cost basis is unaffected - shares and invested come from the trade log, not from prices.",
]:
    sc(pr, f"A{r}", txt, font=NOTE)
    r += 1


# =========================================================== 3. Transactions
tx = wb.create_sheet("Transactions")
tx.sheet_properties.tabColor = "808080"
tx.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMNOPQ",
                  [12, 10, 9, 13, 15, 11, 15, 14, 11, 14, 13, 13, 15, 15, 16, 15, 34]):
    tx.column_dimensions[col].width = w

banner(tx, 1, "MASTER TRADE LOG  —  every buy and sell goes here once", last_col="Q")
sc(tx, "A2", "Shares held and average cost on every holding tab are calculated from this "
             "sheet. Add new trades in the next empty row; keep the formulas in G to P.", font=NOTE)

THEAD = 3
headers(tx, THEAD, ["Date", "Ticker", "Type", "Shares", "Price/share (native)",
                    "FX (GBP/USD)", "Price/share (GBP)", "Gross (GBP)", "Fees (GBP)",
                    "Net cash (GBP)", "Signed shares", "Shares before",
                    "Cost basis before", "Avg cost before", "Cost basis impact",
                    "Realised P/L (GBP)", "Notes"])
tx.freeze_panes = f"A{THEAD+1}"

TFIRST = THEAD + 1
TLAST = TFIRST + 59

dv = DataValidation(type="list", formula1='"BUY,SELL"', allow_blank=True, showDropDown=False)
tx.add_data_validation(dv)
dv.add(f"C{TFIRST}:C{TLAST}")

seed = [(h["tkr"], h["shares"], round(h["invested"] / h["shares"] * GBPUSD, 6))
        for h in HOLDINGS]

for i, row in enumerate(range(TFIRST, TLAST + 1)):
    is_seed = i < len(seed)
    shade = LGREY if (i % 2 == 1) else None
    sc(tx, f"A{row}", PRICE_DATE if is_seed else None, font=BLUE, fill=shade,
       border=True, align="center")
    sc(tx, f"B{row}", seed[i][0] if is_seed else None, font=BLUE, fill=shade,
       border=True, align="center")
    sc(tx, f"C{row}", "BUY" if is_seed else None, font=BLUE, fill=shade,
       border=True, align="center")
    sc(tx, f"D{row}", seed[i][1] if is_seed else None, font=BLUE, fill=shade,
       fmt=SHR, border=True)
    sc(tx, f"E{row}", seed[i][2] if is_seed else None, font=BLUE, fill=shade,
       fmt=USD, border=True)
    sc(tx, f"F{row}", GBPUSD if is_seed else None, font=BLUE, fill=shade,
       fmt='0.00000', border=True)
    sc(tx, f"I{row}", 0 if is_seed else None, font=BLUE, fill=shade, fmt=GBP, border=True)
    sc(tx, f"Q{row}", "Opening position from the Trading 212 Never Sell pie export, "
                      "4 Aug 2026. Replace with your real trade history."
       if is_seed else None, font=BLUE if is_seed else BLACK, fill=shade, border=True)

    g = f'IF($B{row}="","",'
    sc(tx, f"G{row}", f'={g}IFERROR($E{row}/$F{row},0))', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"H{row}", f'={g}$D{row}*$G{row})', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"J{row}", f'={g}IF($C{row}="BUY",$H{row}+$I{row},$H{row}-$I{row}))',
       font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"K{row}", f'={g}IF($C{row}="BUY",$D{row},-$D{row}))',
       font=BLACK, fill=shade, fmt=SHR, border=True)
    if row == TFIRST:
        ps, pc = "0", "0"
    else:
        ps = f'SUMIFS($K${TFIRST}:$K{row-1},$B${TFIRST}:$B{row-1},$B{row})'
        pc = f'SUMIFS($O${TFIRST}:$O{row-1},$B${TFIRST}:$B{row-1},$B{row})'
    sc(tx, f"L{row}", f'={g}{ps})', font=BLACK, fill=shade, fmt=SHR, border=True)
    sc(tx, f"M{row}", f'={g}{pc})', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"N{row}", f'={g}IFERROR($M{row}/$L{row},0))', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"O{row}", f'={g}IF($C{row}="BUY",$J{row},-$D{row}*$N{row}))',
       font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(tx, f"P{row}", f'={g}IF($C{row}="SELL",$J{row}-$D{row}*$N{row},0))',
       font=BLACK, fill=shade, fmt=GBP, border=True)

tx[f"E{TFIRST}"].comment = Comment(
    "Derived so the opening row reproduces your export's cost basis exactly:\n"
    "invested GBP / shares x GBP/USD 1.34499.\n"
    "Source: Never Sell pie export, 4 Aug 2026.", "Portfolio Tracker")

r = TLAST + 2
for txt in [
    "Average-cost accounting: a SELL is booked against the average cost of the shares held "
    "before it, and the difference lands in Realised P/L.",
    "For a GBP-quoted holding (LGEN, BA, SHEL) put the price in pounds in column E and set FX to 1.",
    "Need more than 60 trades? Copy the last row down - every formula is relative.",
]:
    sc(tx, f"A{r}", txt, font=NOTE)
    r += 1


# ============================================================== 4. Dividends
dvs = wb.create_sheet("Dividends")
dvs.sheet_properties.tabColor = "BF8F00"
dvs.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHI", [12, 10, 17, 14, 12, 15, 18, 15, 34]):
    dvs.column_dimensions[col].width = w

banner(dvs, 1, "DIVIDEND LOG  —  feeds the income block on each holding tab", last_col="I")
sc(dvs, "A2", "Log each payment as it lands. Net GBP is what actually hit the account.", font=NOTE)

DHEAD = 3
headers(dvs, DHEAD, ["Date", "Ticker", "Shares at pay date", "Gross (native)",
                     "FX (GBP/USD)", "Gross (GBP)", "Withholding tax (GBP)",
                     "Net (GBP)", "Notes"])
dvs.freeze_panes = f"A{DHEAD+1}"
DFIRST = DHEAD + 1
DLAST = DFIRST + 39
for i, row in enumerate(range(DFIRST, DLAST + 1)):
    shade = LGREY if (i % 2 == 1) else None
    for col, fmt in (("A", None), ("B", None), ("C", SHR), ("D", USD),
                     ("E", '0.00000'), ("G", GBP)):
        sc(dvs, f"{col}{row}", None, font=BLUE, fill=shade, fmt=fmt, border=True)
    g = f'IF($B{row}="","",'
    sc(dvs, f"F{row}", f'={g}IFERROR($D{row}/$E{row},0))', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(dvs, f"H{row}", f'={g}$F{row}-$G{row})', font=BLACK, fill=shade, fmt=GBP, border=True)
    sc(dvs, f"I{row}", None, font=BLUE, fill=shade, border=True)

r = DLAST + 2
for txt in [
    "US dividends in a UK ISA are normally taxed at 15% at source with a W-8BEN on file. "
    "Put that amount in column G.",
    "Of AMZN, MSFT and UBER, only MSFT currently pays a dividend.",
]:
    sc(dvs, f"A{r}", txt, font=NOTE)
    r += 1


# ============================================================= 5. Sell_Rules
sr = wb.create_sheet("Sell_Rules")
sr.sheet_properties.tabColor = "C00000"
sr.sheet_view.showGridLines = False
for col, w in zip("ABCDEF", [40, 14, 64, 14, 14, 14]):
    sr.column_dimensions[col].width = w

banner(sr, 1, "DEFAULT SELL RULES  —  every holding tab inherits these")
SR = {}
r = 3
section(sr, r, "TRIGGER LEVELS"); r += 1
for col, h in zip("ABC", ["Rule", "Level", "What it means"]):
    sc(sr, f"{col}{r}", h, font=HEADF, fill=SLATE, border=True)
r += 1
for key, label, val, meaning in [
    ("take_profit", "Take-profit target (% above avg cost)", 0.50,
     "Full exit candidate. Price >= avg cost x 1.50."),
    ("trim1", "Trim level 1 (% above avg cost)", 0.25,
     "First trim. Take some off the table, let the rest run."),
    ("trim2", "Trim level 2 (% above avg cost)", 0.40,
     "Second trim, closer to the full target."),
    ("stop", "Stop loss (% below avg cost)", 0.20,
     "Hard exit. Price <= avg cost x 0.80."),
    ("trail", "Trailing stop (% below 52-week high)", 0.25,
     "Protects a run-up. Price <= 52w high x 0.75."),
]:
    sc(sr, f"A{r}", label, font=LABEL, border=True)
    sc(sr, f"B{r}", val, font=BLUE, fill=YELLOW, fmt=PCT, border=True, align="center")
    sc(sr, f"C{r}", meaning, font=LABEL, border=True)
    SR[key] = f"Sell_Rules!$B${r}"
    r += 1
r += 1
sc(sr, f"A{r}", "These are placeholders, not advice - set them to your own discipline.", font=NOTE)
r += 1
sc(sr, f"A{r}", "A holding tab can override any level: just type a number over the green link.", font=NOTE)
r += 2
section(sr, r, "SIGNAL PRIORITY  —  the order each holding tab tests"); r += 1
for txt in [
    "1.  No position  ->  NO POSITION",
    "2.  Manual sell price set and reached  ->  SELL - MANUAL TARGET",
    "3.  Price at or below stop loss  ->  SELL - STOP LOSS",
    "4.  Price at or above take-profit  ->  SELL - TARGET HIT",
    "5.  Price at or above trim 2  ->  TRIM 2",
    "6.  Price at or above trim 1  ->  TRIM 1",
    "7.  Price at or below trailing stop  ->  REVIEW - TRAILING STOP",
    "8.  Otherwise  ->  HOLD",
]:
    sc(sr, f"A{r}", txt, font=LABEL)
    r += 1
r += 1
sc(sr, f"A{r}", "Stop loss is tested before take-profit so a crash is never masked by a stale target.",
   font=NOTE)


# =========================================================== 6. Holding tabs
HOLDROW = {}
DASH_MV_ROW = 6      # Dashboard row carrying total market value (asserted later)


def TXR(col):
    return f"Transactions!${col}${TFIRST}:${col}${TLAST}"


def DVR(col):
    return f"Dividends!${col}${DFIRST}:${col}${DLAST}"


def build_holding(h):
    tkr = h["tkr"]
    sh = wb.create_sheet(tkr)
    sh.sheet_properties.tabColor = "2E75B6"
    sh.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", [38, 18, 60, 14, 14, 14]):
        sh.column_dimensions[col].width = w

    banner(sh, 1, f"{tkr}  —  {h['name']}")
    R = {}
    r = 3
    px, er = PROW[tkr], EROW[tkr]

    def line(key, label, formula, font=BLACK, fmt=None, note=None,
             comment=None, fill=None, lab_font=LABEL):
        nonlocal r
        sc(sh, f"A{r}", label, font=lab_font, border=True)
        sc(sh, f"B{r}", formula, font=font, fmt=fmt, border=True,
           align="right" if fmt else "center", comment=comment, fill=fill)
        if note:
            sc(sh, f"C{r}", note, font=NOTE)
        R[key] = r
        r += 1

    # ---- 1. POSITION
    section(sh, r, "1.  POSITION"); r += 1
    line("ticker", "Ticker", tkr, font=BOLD)
    line("company", "Company", h["name"], font=LABEL)
    line("sleeve", "Sleeve / pie", h["sleeve"], font=BLUE)
    line("ccy", "Quote currency", f"=Prices!$C${px}", font=GREEN)
    line("shares", "Shares held", f'=SUMIFS({TXR("K")},{TXR("B")},$B${R["ticker"]})',
         fmt=SHR, note="calculated from the Transactions log - never typed")
    line("invested", "Total invested (GBP)",
         f'=SUMIFS({TXR("O")},{TXR("B")},$B${R["ticker"]})', fmt=GBP)
    line("avgcost", "Average cost per share (GBP)",
         f'=IFERROR($B${R["invested"]}/$B${R["shares"]},0)', fmt=GBP)
    line("pxusd", "Live price (USD)", f"=Prices!$D${px}", font=GREEN, fmt=USD,
         note="update on the Prices tab")
    line("fx", "GBP/USD", f"=Prices!$B${FX_ROW}", font=GREEN, fmt='0.00000')
    line("pxgbp", "Price per share (GBP)", f"=Prices!$K${px}", font=GREEN, fmt=GBP)
    line("mv", "Market value (GBP)", f'=$B${R["shares"]}*$B${R["pxgbp"]}', fmt=GBP)
    line("upl", "Unrealised P/L (GBP)", f'=$B${R["mv"]}-$B${R["invested"]}', fmt=GBP)
    line("uplpct", "Unrealised P/L %", f'=IFERROR($B${R["upl"]}/$B${R["invested"]},0)', fmt=PCT)
    line("daychg", "Day change (GBP)",
         f'=IFERROR($B${R["shares"]}*Prices!$F${px}/$B${R["fx"]},0)', fmt=GBP)
    line("daypct", "Day change %", f'=IFERROR(Prices!$G${px},0)', fmt=PCT2)
    line("weight", "% of portfolio", None, fmt=PCT, note="share of total market value")
    line("rpl", "Realised P/L (GBP)", f'=SUMIFS({TXR("P")},{TXR("B")},$B${R["ticker"]})', fmt=GBP)
    line("divrec", "Dividends received (GBP)",
         f'=SUMIFS({DVR("H")},{DVR("B")},$B${R["ticker"]})', fmt=GBP)
    line("totret", "Total return (GBP)",
         f'=$B${R["upl"]}+$B${R["rpl"]}+$B${R["divrec"]}', fmt=GBP)
    line("totretp", "Total return %", f'=IFERROR($B${R["totret"]}/$B${R["invested"]},0)', fmt=PCT)
    sh[f'B{R["weight"]}'] = f'=IFERROR($B${R["mv"]}/Dashboard!$B${DASH_MV_ROW},0)'
    r += 1

    # ---- 2. VALUATION AND EPS
    section(sh, r, "2.  VALUATION AND EPS"); r += 1
    line("lastfy", f"Last reported FY", f"=Estimates!$D${er}", font=GREEN)
    line("epslast", "EPS last reported FY (USD)", f"=Estimates!$E${er}", font=GREEN, fmt=USD)
    line("eps1", f"EPS {Y1}E  (Year 1)", f"=Estimates!$G${er}", font=GREEN, fmt=USD)
    line("eps2", f"EPS {Y2}E  (Year 2)", f"=Estimates!$H${er}", font=GREEN, fmt=USD)
    line("eps3", f"EPS {Y3}E  (Year 3)", f"=Estimates!$I${er}", font=GREEN, fmt=USD)
    line("epscagr", "EPS CAGR, Year 1 to Year 3", f"=Estimates!$K${er}", font=GREEN, fmt=PCT)
    line("petrail", "P/E on last reported EPS",
         f'=IFERROR($B${R["pxusd"]}/$B${R["epslast"]},0)', fmt=MULT)
    line("pe1", "Forward P/E on Year 1 EPS",
         f'=IFERROR($B${R["pxusd"]}/$B${R["eps1"]},0)', fmt=MULT)
    line("pe3", "P/E on Year 3 EPS",
         f'=IFERROR($B${R["pxusd"]}/$B${R["eps3"]},0)', fmt=MULT,
         note="what you are paying for earnings three years out")
    line("peg", "PEG (fwd P/E / EPS growth)",
         f'=IFERROR($B${R["pe1"]}/($B${R["epscagr"]}*100),0)', fmt=NUM2,
         note="under 1.0 is the classic cheap-for-its-growth marker")
    line("antgt", "Analyst target, consensus (USD)", f"=Estimates!$L${er}", font=GREEN, fmt=USD)
    line("anup", "Upside to analyst target",
         f'=IFERROR($B${R["antgt"]}/$B${R["pxusd"]}-1,0)', fmt=PCT)
    line("anlo", "Analyst target, low / high (USD)",
         f'=Estimates!$N${er}&"  to  "&Estimates!$O${er}', font=GREEN)
    line("hi52", "52-week high (USD)", f"=Prices!$H${px}", font=GREEN, fmt=USD)
    line("lo52", "52-week low (USD)", f"=Prices!$I${px}", font=GREEN, fmt=USD)
    line("offhi", "% below 52-week high", f"=IFERROR(Prices!$J${px},0)", fmt=PCT)
    r += 1

    # ---- 3. YOUR 3-YEAR PRICE TARGETS
    section(sh, r, "3.  YOUR 3-YEAR PRICE TARGETS  —  what you are actually holding for"); r += 1
    line("tgtpe", "Your target P/E", f"=Estimates!$P${er}", font=BLUE, fmt=MULT, fill=YELLOW,
         note="starts at the multiple the street's target implies - overwrite with your own",
         comment="This is YOUR multiple. It starts at consensus target / Year 1 EPS so the "
                 "Year 1 target matches the street, then diverges as you set your own view.")
    line("oeps1", f"Your EPS {Y1} (Year 1)", f"=Estimates!$G${er}", font=BLUE, fmt=USD, fill=YELLOW,
         note="starts at consensus - change it if you disagree")
    line("oeps2", f"Your EPS {Y2} (Year 2)", f"=Estimates!$H${er}", font=BLUE, fmt=USD, fill=YELLOW)
    line("oeps3", f"Your EPS {Y3} (Year 3)", f"=Estimates!$I${er}", font=BLUE, fmt=USD, fill=YELLOW)
    line("t1usd", f"Target price {Y1} (USD)",
         f'=IFERROR($B${R["tgtpe"]}*$B${R["oeps1"]},0)', fmt=USD, lab_font=BOLD)
    line("t2usd", f"Target price {Y2} (USD)",
         f'=IFERROR($B${R["tgtpe"]}*$B${R["oeps2"]},0)', fmt=USD, lab_font=BOLD)
    line("t3usd", f"Target price {Y3} (USD)",
         f'=IFERROR($B${R["tgtpe"]}*$B${R["oeps3"]},0)', fmt=USD, lab_font=BOLD)
    line("t1gbp", f"Target price {Y1} (GBP)", f'=IFERROR($B${R["t1usd"]}/$B${R["fx"]},0)', fmt=GBP)
    line("t2gbp", f"Target price {Y2} (GBP)", f'=IFERROR($B${R["t2usd"]}/$B${R["fx"]},0)', fmt=GBP)
    line("t3gbp", f"Target price {Y3} (GBP)", f'=IFERROR($B${R["t3usd"]}/$B${R["fx"]},0)', fmt=GBP)
    line("up1", "Upside to Year 1 target", f'=IFERROR($B${R["t1usd"]}/$B${R["pxusd"]}-1,0)', fmt=PCT)
    line("up2", "Upside to Year 2 target", f'=IFERROR($B${R["t2usd"]}/$B${R["pxusd"]}-1,0)', fmt=PCT)
    line("up3", "Upside to Year 3 target", f'=IFERROR($B${R["t3usd"]}/$B${R["pxusd"]}-1,0)', fmt=PCT)
    line("cagr", "Required annual return to Year 3",
         f'=IFERROR(($B${R["t3usd"]}/$B${R["pxusd"]})^(1/3)-1,0)', fmt=PCT,
         note="the compound return this holding must deliver for your target to land")
    line("v1", f"Your holding worth at {Y1} (GBP)",
         f'=$B${R["shares"]}*$B${R["t1gbp"]}', fmt=GBP)
    line("v2", f"Your holding worth at {Y2} (GBP)",
         f'=$B${R["shares"]}*$B${R["t2gbp"]}', fmt=GBP)
    line("v3", f"Your holding worth at {Y3} (GBP)",
         f'=$B${R["shares"]}*$B${R["t3gbp"]}', fmt=GBP, lab_font=BOLD)
    line("gain3", "Gain from here to Year 3 (GBP)", f'=$B${R["v3"]}-$B${R["mv"]}', fmt=GBP)
    line("prog", "Progress from cost to Year 3 target",
         f'=IFERROR(($B${R["pxgbp"]}-$B${R["avgcost"]})/($B${R["t3gbp"]}-$B${R["avgcost"]}),0)',
         fmt=PCT, note="0% = at your cost, 100% = target reached")
    sc(sh, f"A{r}", "TARGET STATUS", font=BOLD, fill=GOLD, border=True)
    sc(sh, f"B{r}",
       f'=IF($B${R["shares"]}=0,"NO POSITION",'
       f'IF($B${R["pxgbp"]}>=$B${R["t3gbp"]},"YEAR 3 TARGET MET",'
       f'IF($B${R["pxgbp"]}>=$B${R["t2gbp"]},"AHEAD - PAST YEAR 2",'
       f'IF($B${R["pxgbp"]}>=$B${R["t1gbp"]},"AHEAD - PAST YEAR 1",'
       f'IF($B${R["pxgbp"]}>=$B${R["avgcost"]},"ON TRACK","BEHIND - BELOW COST")))))',
       font=BIG, fill=GOLD, border=True, align="center")
    sc(sh, f"C{r}", "where the price sits against your own 3-year path", font=NOTE)
    sh.row_dimensions[r].height = 22
    R["tstatus"] = r
    r += 2

    # ---- 4. SELL DISCIPLINE
    section(sh, r, "4.  SELL DISCIPLINE  —  the exit plan, decided in advance"); r += 1
    line("tp", "Take-profit target %", f"={SR['take_profit']}", font=GREEN, fmt=PCT,
         note="inherits Sell_Rules - type over it to override")
    line("tpp", "Take-profit price (GBP)", f'=$B${R["avgcost"]}*(1+$B${R["tp"]})', fmt=GBP)
    line("tr1", "Trim level 1 %", f"={SR['trim1']}", font=GREEN, fmt=PCT)
    line("tr1p", "Trim 1 price (GBP)", f'=$B${R["avgcost"]}*(1+$B${R["tr1"]})', fmt=GBP)
    line("tr2", "Trim level 2 %", f"={SR['trim2']}", font=GREEN, fmt=PCT)
    line("tr2p", "Trim 2 price (GBP)", f'=$B${R["avgcost"]}*(1+$B${R["tr2"]})', fmt=GBP)
    line("sl", "Stop loss %", f"={SR['stop']}", font=GREEN, fmt=PCT)
    line("slp", "Stop-loss price (GBP)", f'=$B${R["avgcost"]}*(1-$B${R["sl"]})', fmt=GBP)
    line("trl", "Trailing stop % off 52w high", f"={SR['trail']}", font=GREEN, fmt=PCT)
    line("trlp", "Trailing stop price (GBP)",
         f'=IFERROR($B${R["hi52"]}/$B${R["fx"]}*(1-$B${R["trl"]}),0)', fmt=GBP)
    line("manual", "Manual sell price (GBP)", None, font=BLUE, fmt=GBP, fill=YELLOW,
         note="optional hard number that overrides everything above",
         comment="Leave blank to use the percentage rules. Set a price here and it takes priority.")
    sc(sh, f"A{r}", "SIGNAL", font=BOLD, fill=BANNER, border=True)
    sc(sh, f"B{r}",
       f'=IF($B${R["shares"]}=0,"NO POSITION",'
       f'IF(AND($B${R["manual"]}>0,$B${R["pxgbp"]}>=$B${R["manual"]}),"SELL - MANUAL TARGET",'
       f'IF($B${R["pxgbp"]}<=$B${R["slp"]},"SELL - STOP LOSS",'
       f'IF($B${R["pxgbp"]}>=$B${R["tpp"]},"SELL - TARGET HIT",'
       f'IF($B${R["pxgbp"]}>=$B${R["tr2p"]},"TRIM 2",'
       f'IF($B${R["pxgbp"]}>=$B${R["tr1p"]},"TRIM 1",'
       f'IF(AND($B${R["trlp"]}>0,$B${R["pxgbp"]}<=$B${R["trlp"]}),"REVIEW - TRAILING STOP",'
       f'"HOLD")))))))',
       font=BIG, fill=BANNER, border=True, align="center")
    sc(sh, f"C{r}", "recalculated the moment the price changes", font=NOTE)
    sh.row_dimensions[r].height = 22
    R["signal"] = r
    r += 1
    line("selval", "Value if sold today (GBP)", f'=$B${R["mv"]}', fmt=GBP)
    line("selfee", "Estimated selling costs (GBP)",
         f'=ROUND($B${R["mv"]}*0.0015,2)', fmt=GBP,
         note="Trading 212 FX fee, 0.15% of the trade",
         comment="0.15% of market value - Trading 212's FX conversion fee on a USD sale. "
                 "Overwrite if your fee differs.")
    line("selnet", "Net proceeds if sold (GBP)", f'=$B${R["selval"]}-$B${R["selfee"]}', fmt=GBP)
    line("sellock", "Gain locked in if sold (GBP)", f'=$B${R["selnet"]}-$B${R["invested"]}', fmt=GBP)
    line("sellockp", "Gain locked in %", f'=IFERROR($B${R["sellock"]}/$B${R["invested"]},0)', fmt=PCT)
    r += 1

    # ---- 5. DIVIDENDS
    section(sh, r, "5.  DIVIDEND INCOME"); r += 1
    line("dps", "Annual dividend per share (USD)", h["dps"], font=BLUE, fmt=USD, fill=YELLOW,
         note="0 if it does not pay one",
         comment="Declared annual dividend per share, USD.\n"
                 "MSFT pays; AMZN and UBER currently do not.")
    line("dyield", "Dividend yield on price", f'=IFERROR($B${R["dps"]}/$B${R["pxusd"]},0)', fmt=PCT2)
    line("dyoc", "Yield on your cost",
         f'=IFERROR($B${R["dps"]}/$B${R["fx"]}/$B${R["avgcost"]},0)', fmt=PCT2,
         note="the number that matters once you are in profit")
    line("dinc", "Estimated annual income (GBP)",
         f'=IFERROR($B${R["shares"]}*$B${R["dps"]}/$B${R["fx"]},0)', fmt=GBP)
    line("dgot", "Dividends received to date (GBP)", f'=$B${R["divrec"]}', fmt=GBP)
    r += 1

    # ---- 6. TRADE HISTORY
    section(sh, r, "6.  TRADE HISTORY  —  summarised from the Transactions log"); r += 1
    t = f'{TXR("B")},$B${R["ticker"]}'
    line("ntrades", "Number of trades", f'=COUNTIFS({TXR("B")},$B${R["ticker"]})', fmt=INT)
    line("bought", "Shares bought", f'=SUMIFS({TXR("D")},{t},{TXR("C")},"BUY")', fmt=SHR)
    line("sold", "Shares sold", f'=SUMIFS({TXR("D")},{t},{TXR("C")},"SELL")', fmt=SHR)
    line("cashin", "Total cash invested (GBP)", f'=SUMIFS({TXR("J")},{t},{TXR("C")},"BUY")', fmt=GBP)
    line("cashout", "Total cash from sales (GBP)", f'=SUMIFS({TXR("J")},{t},{TXR("C")},"SELL")', fmt=GBP)
    line("feespaid", "Total fees paid (GBP)", f'=SUMIFS({TXR("I")},{t})', fmt=GBP)
    r += 1

    # ---- 7. THESIS
    section(sh, r, "7.  THESIS AND NOTES"); r += 1
    for label, hint in [
        ("Why I own it", "the one-line reason - if you cannot write it, you should not hold it"),
        ("What would make me sell", "name the broken thing, not a price"),
        ("Key risk", ""),
        ("Next catalyst / earnings date", ""),
        ("Last reviewed", ""),
    ]:
        sc(sh, f"A{r}", label, font=LABEL, border=True)
        sc(sh, f"B{r}", None, font=BLUE, fill=YELLOW, border=True)
        sh.merge_cells(f"B{r}:C{r}")
        sc(sh, f"C{r}", None, fill=YELLOW, border=True)
        if hint:
            sc(sh, f"D{r}", hint, font=NOTE)
        r += 1

    # conditional formatting
    for key in ("upl", "uplpct", "totret", "totretp", "rpl", "daychg", "daypct",
                "sellock", "sellockp", "gain3", "anup", "up1", "up2", "up3"):
        rr = R[key]
        sh.conditional_formatting.add(f"B{rr}", CellIsRule(operator="lessThan", formula=["0"], font=red_f))
        sh.conditional_formatting.add(f"B{rr}", CellIsRule(operator="greaterThan", formula=["0"], font=green_f))

    sig = f'B{R["signal"]}'
    for pat, fillc, fontc in (("SELL", "FFC7CE", "9C0006"), ("TRIM", "FFEB9C", "9C6500"),
                              ("REVIEW", "FFEB9C", "9C6500")):
        sh.conditional_formatting.add(sig, FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{pat}",{sig}))'],
            fill=PatternFill("solid", fgColor=fillc),
            font=Font(name=FONT, size=12, bold=True, color=fontc)))
    sh.conditional_formatting.add(sig, FormulaRule(
        formula=[f'EXACT({sig},"HOLD")'], fill=PatternFill("solid", fgColor="C6EFCE"),
        font=Font(name=FONT, size=12, bold=True, color="006100")))

    ts = f'B{R["tstatus"]}'
    for pat, fillc, fontc in (("AHEAD", "C6EFCE", "006100"), ("MET", "C6EFCE", "006100"),
                              ("BEHIND", "FFC7CE", "9C0006")):
        sh.conditional_formatting.add(ts, FormulaRule(
            formula=[f'ISNUMBER(SEARCH("{pat}",{ts}))'],
            fill=PatternFill("solid", fgColor=fillc),
            font=Font(name=FONT, size=12, bold=True, color=fontc)))

    sh.conditional_formatting.add(f"B{R['prog']}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="638EC6"))

    HOLDROW[tkr] = R


for h in HOLDINGS:
    build_holding(h)


# ================================================================ 7. Targets
tg = wb.create_sheet("Targets")
tg.sheet_properties.tabColor = "ED7D31"
tg.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJ", [26, 15, 15, 15, 15, 15, 15, 15, 15, 20]):
    tg.column_dimensions[col].width = w

banner(tg, 1, "3-YEAR TARGETS  —  where this portfolio is heading and whether it is on pace",
       last_col="J")
sc(tg, "A2", f"Targets come from each holding tab: your own target P/E x your own EPS. "
             f"Year 1 = {Y1}, Year 2 = {Y2}, Year 3 = {Y3}.", font=NOTE)

r = 4
section(tg, r, "PROJECTED VALUE BY HOLDING", last_col="J"); r += 1
GHEAD = r
headers(tg, GHEAD, ["Ticker", "Shares", "Price now (GBP)", f"{Y1} target",
                    f"{Y2} target", f"{Y3} target", "Value now",
                    f"Value at {Y3}", f"Gain to {Y3}", "Required return p.a."])
r += 1
GFIRST = r
for h in HOLDINGS:
    tkr = h["tkr"]
    R = HOLDROW[tkr]
    for col, f, fmt in (
        ("A", f"='{tkr}'!$B${R['ticker']}", None),
        ("B", f"='{tkr}'!$B${R['shares']}", SHR),
        ("C", f"='{tkr}'!$B${R['pxgbp']}", GBP),
        ("D", f"='{tkr}'!$B${R['t1gbp']}", GBP),
        ("E", f"='{tkr}'!$B${R['t2gbp']}", GBP),
        ("F", f"='{tkr}'!$B${R['t3gbp']}", GBP),
        ("G", f"='{tkr}'!$B${R['mv']}", GBP),
        ("H", f"='{tkr}'!$B${R['v3']}", GBP),
        ("I", f"='{tkr}'!$B${R['gain3']}", GBP),
        ("J", f"='{tkr}'!$B${R['cagr']}", PCT),
    ):
        sc(tg, f"{col}{r}", f, font=BOLD if col == "A" else GREEN, fmt=fmt,
           border=True, align="center" if col == "A" else "right")
    r += 1
GLAST = r - 1
sc(tg, f"A{r}", "TOTAL", font=BOLD, fill=BANNER, border=True)
for col in "BCDEF":
    sc(tg, f"{col}{r}", None, fill=BANNER, border=True)
for col in "GHI":
    sc(tg, f"{col}{r}", f"=SUM({col}{GFIRST}:{col}{GLAST})", font=BOLD, fill=BANNER,
       fmt=GBP, border=True, align="right")
sc(tg, f"J{r}", f"=IFERROR((H{r}/G{r})^(1/3)-1,0)", font=BOLD, fill=BANNER, fmt=PCT,
   border=True, align="right")
GTOT = r
r += 2

section(tg, r, "PORTFOLIO PATH", last_col="J"); r += 1
PATH = {}
for key, label, formula, fmt in [
    ("now", "Value today (GBP)", f"=G{GTOT}", GBP0),
    ("v1", f"Projected value {Y1} (GBP)",
     "=" + "+".join([f"'{h['tkr']}'!$B${HOLDROW[h['tkr']]['v1']}" for h in HOLDINGS]), GBP0),
    ("v2", f"Projected value {Y2} (GBP)",
     "=" + "+".join([f"'{h['tkr']}'!$B${HOLDROW[h['tkr']]['v2']}" for h in HOLDINGS]), GBP0),
    ("v3", f"Projected value {Y3} (GBP)", f"=H{GTOT}", GBP0),
    ("gain", f"Total gain to {Y3} (GBP)", f"=I{GTOT}", GBP0),
    ("cagr", f"Portfolio return needed p.a.", f"=J{GTOT}", PCT),
]:
    sc(tg, f"A{r}", label, font=BOLD, fill=LGREY, border=True)
    sc(tg, f"B{r}", formula, font=BLACK, fmt=fmt, border=True, align="right")
    PATH[key] = r
    r += 1
r += 1

section(tg, r, "THE £1,000,000 GOAL", last_col="J"); r += 1
G = {}
for key, label, formula, fmt, font, fill in [
    ("goal", "Target portfolio value (GBP)", GOAL, GBP0, BLUE, YELLOW),
    ("contrib", "Planned monthly contribution (GBP)", 0, GBP0, BLUE, YELLOW),
    ("now", "Value today (GBP)", f"=B{PATH['now']}", GBP0, GREEN, None),
    ("pct", "% of goal reached today", None, PCT, BLACK, None),
    ("gap", "Gap to goal (GBP)", None, GBP0, BLACK, None),
    ("v3", f"Projected value at {Y3} (GBP)", f"=B{PATH['v3']}", GBP0, GREEN, None),
    ("pct3", f"% of goal at {Y3}", None, PCT, BLACK, None),
    ("yrs", "Years to goal at that return, no new money", None, NUM2, BLACK, None),
    ("need", f"Monthly contribution to hit goal by {Y3}", None, GBP0, BLACK, None),
    ("withc", f"Projected {Y3} value incl. your contributions", None, GBP0, BLACK, None),
]:
    sc(tg, f"A{r}", label, font=BOLD, fill=LGREY if fill is None else None, border=True)
    sc(tg, f"B{r}", formula, font=font, fmt=fmt, border=True, align="right", fill=fill)
    G[key] = r
    r += 1

tg[f"B{G['pct']}"] = f"=IFERROR(B{G['now']}/B{G['goal']},0)"
tg[f"B{G['gap']}"] = f"=B{G['goal']}-B{G['now']}"
tg[f"B{G['pct3']}"] = f"=IFERROR(B{G['v3']}/B{G['goal']},0)"
tg[f"B{G['yrs']}"] = (f"=IFERROR(LN(B{G['goal']}/B{G['now']})"
                      f"/LN(1+B{PATH['cagr']}),0)")
# monthly annuity needed: (goal - now*(1+r)^n) / (((1+r)^n - 1)/r), r = cagr/12, n = 36
_r = f"(B{PATH['cagr']}/12)"
_n = "36"
tg[f"B{G['need']}"] = (f"=IFERROR(MAX(0,(B{G['goal']}-B{G['now']}*(1+{_r})^{_n})"
                       f"/(((1+{_r})^{_n}-1)/{_r})),0)")
tg[f"B{G['withc']}"] = (f"=IFERROR(B{G['now']}*(1+{_r})^{_n}"
                        f"+B{G['contrib']}*(((1+{_r})^{_n}-1)/{_r}),0)")

tg[f"B{G['goal']}"].comment = Comment(
    "Your target. Change it and every line below re-solves.", "Portfolio Tracker")
tg[f"B{G['contrib']}"].comment = Comment(
    "What you actually plan to add each month. The last line shows where that lands you.",
    "Portfolio Tracker")
tg[f"B{G['need']}"].comment = Comment(
    "Solves the monthly payment that closes the gap by Year 3, assuming the portfolio "
    "compounds at the required return above. Zero means the targets alone get you there.",
    "Portfolio Tracker")

r += 1
for txt in [
    "This is arithmetic on your own assumptions, not a forecast. Change the target P/E or EPS "
    "on any holding tab and every number here moves.",
    "Three holdings will not reach £1m on their own - this tab is here to show the gap honestly "
    "as the other sleeves get added.",
]:
    sc(tg, f"A{r}", txt, font=NOTE)
    r += 1

for key in ("gain",):
    tg.conditional_formatting.add(f"B{PATH[key]}",
        CellIsRule(operator="lessThan", formula=["0"], font=red_f))
    tg.conditional_formatting.add(f"B{PATH[key]}",
        CellIsRule(operator="greaterThan", formula=["0"], font=green_f))
tg.conditional_formatting.add(f"I{GFIRST}:I{GLAST}",
    CellIsRule(operator="lessThan", formula=["0"], font=red_f))
tg.conditional_formatting.add(f"I{GFIRST}:I{GLAST}",
    CellIsRule(operator="greaterThan", formula=["0"], font=green_f))
tg.conditional_formatting.add(f"B{G['pct']}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="ED7D31"))
tg.conditional_formatting.add(f"B{G['pct3']}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="ED7D31"))


# ============================================================== 8. Dashboard
db = wb.create_sheet("Dashboard", 0)
db.sheet_properties.tabColor = "FF0000"
db.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGHIJKLMN",
                  [30, 17, 24, 13, 13, 13, 14, 14, 14, 10, 10, 13, 20, 20]):
    db.column_dimensions[col].width = w

banner(db, 1, "PORTFOLIO DASHBOARD", last_col="N")
sc(db, "A2", "Prices as at", font=BOLD)
sc(db, "B2", f"=Prices!$B${DATE_ROW}", font=GREEN, align="center")
sc(db, "C2", "update on the Prices tab", font=NOTE)

r = 4
section(db, r, "PORTFOLIO SUMMARY", last_col="N"); r += 1
SUM = {}
for key, label, fmt in [
    ("invested", "Total invested (GBP)", GBP0),
    ("mv", "Current market value (GBP)", GBP0),
    ("upl", "Unrealised P/L (GBP)", GBP0),
    ("uplpct", "Unrealised P/L %", PCT),
    ("rpl", "Realised P/L (GBP)", GBP0),
    ("div", "Dividends received (GBP)", GBP0),
    ("totret", "Total return (GBP)", GBP0),
    ("totretp", "Total return %", PCT),
    ("daychg", "Day change (GBP)", GBP0),
    ("daypct", "Day change %", PCT2),
    ("nhold", "Number of holdings", INT),
    ("biggest", "Largest position", None),
    ("v3", f"Projected value at {Y3} (GBP)", GBP0),
    ("goalpct", "% of £1m goal today", PCT),
]:
    sc(db, f"A{r}", label, font=BOLD, fill=LGREY, border=True)
    sc(db, f"B{r}", None, font=BLACK, fmt=fmt, border=True, align="right")
    SUM[key] = r
    r += 1
assert SUM["mv"] == DASH_MV_ROW, f"market value on row {SUM['mv']}, holding tabs expect {DASH_MV_ROW}"

r += 1
section(db, r, "HOLDINGS", last_col="N"); r += 1
HHEAD = r
headers(db, HHEAD, ["Ticker", "Company", "Sleeve", "Shares", "Avg cost", "Price",
                    "Invested", "Market value", "Unreal. P/L", "P/L %", "Weight",
                    f"{Y3} target", "Signal", "Target status"])
db.freeze_panes = f"A{HHEAD+1}"
r += 1
HFIRST = r
for h in HOLDINGS:
    tkr = h["tkr"]
    R = HOLDROW[tkr]
    for col, f, fmt, font, align in [
        ("A", f"='{tkr}'!$B${R['ticker']}", None, BOLD, "center"),
        ("B", f"='{tkr}'!$B${R['company']}", None, GREEN, None),
        ("C", f"='{tkr}'!$B${R['sleeve']}", None, GREEN, "center"),
        ("D", f"='{tkr}'!$B${R['shares']}", SHR, GREEN, None),
        ("E", f"='{tkr}'!$B${R['avgcost']}", GBP, GREEN, None),
        ("F", f"='{tkr}'!$B${R['pxgbp']}", GBP, GREEN, None),
        ("G", f"='{tkr}'!$B${R['invested']}", GBP, GREEN, None),
        ("H", f"='{tkr}'!$B${R['mv']}", GBP, GREEN, None),
        ("I", f"='{tkr}'!$B${R['upl']}", GBP, GREEN, None),
        ("J", f"='{tkr}'!$B${R['uplpct']}", PCT, GREEN, None),
        ("K", f"='{tkr}'!$B${R['weight']}", PCT, GREEN, None),
        ("L", f"='{tkr}'!$B${R['t3gbp']}", GBP, GREEN, None),
        ("M", f"='{tkr}'!$B${R['signal']}", None, BOLD, "center"),
        ("N", f"='{tkr}'!$B${R['tstatus']}", None, BOLD, "center"),
    ]:
        sc(db, f"{col}{r}", f, font=font, fmt=fmt, border=True, align=align or "right")
    r += 1
HLAST = r - 1

sc(db, f"A{r}", "TOTAL", font=BOLD, fill=BANNER, border=True)
for col in "BCDEF":
    sc(db, f"{col}{r}", None, fill=BANNER, border=True)
for col in "GHI":
    sc(db, f"{col}{r}", f"=SUM({col}{HFIRST}:{col}{HLAST})", font=BOLD, fill=BANNER,
       fmt=GBP, border=True, align="right")
sc(db, f"J{r}", f"=IFERROR(I{r}/G{r},0)", font=BOLD, fill=BANNER, fmt=PCT, border=True, align="right")
sc(db, f"K{r}", f"=SUM(K{HFIRST}:K{HLAST})", font=BOLD, fill=BANNER, fmt=PCT, border=True, align="right")
for col in "LMN":
    sc(db, f"{col}{r}", None, fill=BANNER, border=True)
TOT = r
r += 2
for txt in [
    "Green figures are pulled live from the holding tabs. Change nothing here - edit the holding "
    "tab, the Prices tab or the Estimates tab.",
    "Weight is each position as a share of current market value, not of invested capital.",
    f"{Y3} target is your own target P/E x your own {Y3} EPS, both set on the holding tab.",
]:
    sc(db, f"A{r}", txt, font=NOTE)
    r += 1

db[f"B{SUM['invested']}"] = f"=G{TOT}"
db[f"B{SUM['mv']}"] = f"=H{TOT}"
db[f"B{SUM['upl']}"] = f"=I{TOT}"
db[f"B{SUM['uplpct']}"] = f"=IFERROR(B{SUM['upl']}/B{SUM['invested']},0)"
db[f"B{SUM['rpl']}"] = f"=SUM({TXR('P')})"
db[f"B{SUM['div']}"] = f"=SUM({DVR('H')})"
db[f"B{SUM['totret']}"] = f"=B{SUM['upl']}+B{SUM['rpl']}+B{SUM['div']}"
db[f"B{SUM['totretp']}"] = f"=IFERROR(B{SUM['totret']}/B{SUM['invested']},0)"
db[f"B{SUM['daychg']}"] = "=" + "+".join(
    [f"'{h['tkr']}'!$B${HOLDROW[h['tkr']]['daychg']}" for h in HOLDINGS])
db[f"B{SUM['daypct']}"] = (f"=IFERROR(B{SUM['daychg']}/(B{SUM['mv']}-B{SUM['daychg']}),0)")
db[f"B{SUM['nhold']}"] = f'=COUNTIF(A{HFIRST}:A{HLAST},"?*")'
db[f"B{SUM['biggest']}"] = (f'=IFERROR(INDEX(A{HFIRST}:A{HLAST},'
                            f'MATCH(MAX(H{HFIRST}:H{HLAST}),H{HFIRST}:H{HLAST},0)),"")')
db[f"B{SUM['v3']}"] = f"=Targets!$B${PATH['v3']}"
db[f"B{SUM['goalpct']}"] = f"=Targets!$B${G['pct']}"

for key in ("upl", "uplpct", "rpl", "totret", "totretp", "daychg", "daypct"):
    cc = f"B{SUM[key]}"
    db.conditional_formatting.add(cc, CellIsRule(operator="lessThan", formula=["0"], font=red_f))
    db.conditional_formatting.add(cc, CellIsRule(operator="greaterThan", formula=["0"], font=green_f))
db.conditional_formatting.add(f"I{HFIRST}:J{HLAST}",
    CellIsRule(operator="lessThan", formula=["0"], font=red_f))
db.conditional_formatting.add(f"I{HFIRST}:J{HLAST}",
    CellIsRule(operator="greaterThan", formula=["0"], font=green_f))
for pat, fillc, fontc in (("SELL", "FFC7CE", "9C0006"), ("TRIM", "FFEB9C", "9C6500"),
                          ("HOLD", "C6EFCE", "006100")):
    db.conditional_formatting.add(f"M{HFIRST}:M{HLAST}", FormulaRule(
        formula=[f'ISNUMBER(SEARCH("{pat}",$M{HFIRST}))'],
        fill=PatternFill("solid", fgColor=fillc),
        font=Font(name=FONT, size=10, bold=True, color=fontc)))
for pat, fillc, fontc in (("AHEAD", "C6EFCE", "006100"), ("MET", "C6EFCE", "006100"),
                          ("BEHIND", "FFC7CE", "9C0006")):
    db.conditional_formatting.add(f"N{HFIRST}:N{HLAST}", FormulaRule(
        formula=[f'ISNUMBER(SEARCH("{pat}",$N{HFIRST}))'],
        fill=PatternFill("solid", fgColor=fillc),
        font=Font(name=FONT, size=10, bold=True, color=fontc)))
db.conditional_formatting.add(f"K{HFIRST}:K{HLAST}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="638EC6"))
db.conditional_formatting.add(f"B{SUM['goalpct']}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=1, color="ED7D31"))


# ================================================================= 9. README
rd = wb.create_sheet("README")
rd.sheet_properties.tabColor = "1F3864"
rd.sheet_view.showGridLines = False
for col, w in zip("ABCDEF", [28, 64, 30, 14, 14, 14]):
    rd.column_dimensions[col].width = w

banner(rd, 1, "LIVE PORTFOLIO TRACKER  —  how to use this workbook")
r = 3
sc(rd, f"A{r}", "Built", font=BOLD)
sc(rd, f"B{r}", "4 August 2026. Holdings: AMZN, MSFT, UBER (Never Sell sleeve).", font=LABEL)
r += 2

section(rd, r, "COLOUR LEGEND  —  what you may type into"); r += 1
for name, desc, fill, font in [
    ("Yellow fill, blue text", "TYPE HERE. Prices, EPS, your target P/E, your sell rules.", YELLOW, BLUE),
    ("Blue text, no fill", "Hardcoded input you can change (opening balances, fees).", None, BLUE),
    ("Black text", "Formula. Do not overtype - it breaks the calculation chain.", None, BLACK),
    ("Green text", "Link pulling from another tab in this workbook.", None, GREEN),
]:
    sc(rd, f"A{r}", name, font=font, fill=fill, border=True)
    sc(rd, f"B{r}", desc, font=LABEL, border=True)
    r += 1
r += 1

section(rd, r, "TAB GUIDE"); r += 1
for name, desc in [
    ("Dashboard", "Whole-portfolio totals, every holding side by side, weights, sell signals and target status."),
    ("Targets", "The 3-year plan: projected value each year, required return, and progress toward £1m."),
    ("Prices", "The price engine, plus a reconciliation against your Trading 212 pie export."),
    ("Estimates", "Consensus EPS by fiscal year and analyst price targets. Refresh quarterly."),
    ("AMZN / MSFT / UBER", "Per holding: position, valuation & EPS, your 3-year targets, sell discipline, dividends, trades, thesis."),
    ("Transactions", "Master trade log. Shares and average cost are calculated from it, never typed."),
    ("Dividends", "Dividend log, feeding the income block on each holding tab."),
    ("Sell_Rules", "Default trigger levels, inherited by every holding tab."),
]:
    sc(rd, f"A{r}", name, font=BOLD, border=True)
    sc(rd, f"B{r}", desc, font=LABEL, border=True)
    r += 1
r += 1

section(rd, r, "HOW THE 3-YEAR TARGETS WORK"); r += 1
for name, desc in [
    ("1. EPS", f"Estimates carries consensus EPS for {Y1}, {Y2} and {Y3}. Each holding tab copies "
               "those into three yellow cells you can overwrite with your own view."),
    ("2. Your multiple", "Set your target P/E. It starts at the multiple the consensus price target "
                         "implies on Year 1 earnings, so Year 1 begins level with the street."),
    ("3. Target price", "Target price = your target P/E x your EPS, for each of the three years, in USD and GBP."),
    ("4. Are you on pace", "TARGET STATUS compares today's price against your own path, and the Targets "
                           "tab rolls all three holdings into one portfolio projection."),
]:
    sc(rd, f"A{r}", name, font=BOLD, border=True)
    sc(rd, f"B{r}", desc, font=LABEL, border=True, wrap=True)
    rd.row_dimensions[r].height = 30
    r += 1
r += 1

section(rd, r, "MAKING THE PRICES LIVE  —  three options, pick one"); r += 1
for name, desc in [
    ("1. Stocks data type", "Microsoft 365. On Prices, select A7:A9, then Data > Stocks. Then set D7 "
                            "to  =A7.Price  and fill down. Refresh with Data > Refresh All."),
    ("2. STOCKHISTORY", "Microsoft 365. Put  =STOCKHISTORY(A7,TODAY()-5,TODAY(),0,0,1)  in a spare "
                        "cell and point D7 at its last row. Recalculates on open."),
    ("3. Manual / paste", "Works in every Excel including web and Mac. Type or paste into D7:D9."),
]:
    sc(rd, f"A{r}", name, font=BOLD, border=True)
    sc(rd, f"B{r}", desc, font=LABEL, border=True, wrap=True)
    rd.row_dimensions[r].height = 34
    r += 1
r += 1
sc(rd, f"A{r}", "Note", font=BOLD)
sc(rd, f"B{r}", "A generated .xlsx cannot ship already wired to a market feed - Excel only trusts a "
                "connection you enable yourself. Options 1 and 2 are one-time setup.", font=NOTE, wrap=True)
rd.row_dimensions[r].height = 30
r += 2

section(rd, r, "ASSUMPTIONS AND SOURCES  —  read before trusting a number"); r += 1
for name, desc in [
    ("Shares and cost basis", "Trading 212 Never Sell pie export, Never-sell-2026-08-04T11-15-05.211Z.csv."),
    ("Prices, 52-week range", "FMP end-of-day close, 3 August 2026 - the last close before your export. "
                              "52-week range is on a closing basis, not intraday."),
    ("GBP/USD 1.34499", "FMP forex quote GBPUSD, 4 August 2026."),
    ("EPS estimates", "FMP analyst consensus (epsAvg) by fiscal year, pulled 4 August 2026. "
                      "AMZN and UBER run to December; MSFT's fiscal year ends in June."),
    ("Analyst price targets", "FMP price-target consensus, 4 August 2026."),
    ("Your target P/E", "Seeded at consensus target / Year 1 EPS. It is a starting point, not a recommendation."),
    ("Sell rule defaults", "50% take-profit, 25% / 40% trims, 20% stop, 25% trailing. Placeholders, not advice."),
    ("Selling costs", "0.15% of market value - Trading 212's FX fee on a USD sale."),
    ("Pie export gap", "The export's prices lag the 3 Aug close, so it understates AMZN and MSFT by about "
                       "2%. See the reconciliation block on Prices."),
]:
    sc(rd, f"A{r}", name, font=BOLD, border=True)
    sc(rd, f"B{r}", desc, font=LABEL, border=True, wrap=True)
    rd.row_dimensions[r].height = 30 if len(desc) > 85 else 15
    r += 1
r += 1
sc(rd, f"A{r}", "Not financial advice. Projections are arithmetic on your own assumptions, not forecasts.",
   font=NOTE)

# ----------------------------------------------------------------- finalise
order = ["Dashboard", "Targets", "README", "Prices", "Estimates"] + \
        [h["tkr"] for h in HOLDINGS] + ["Transactions", "Dividends", "Sell_Rules"]
wb._sheets = [wb[n] for n in order]
wb.active = 0

# openpyxl writes formulas with no cached results; make Excel compute on open.
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("written:", OUT)
print("sheets:", wb.sheetnames)
