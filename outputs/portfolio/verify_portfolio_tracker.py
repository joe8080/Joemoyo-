#!/usr/bin/env python3
"""Independent check of the tracker: every reference resolves, and the numbers
Excel will produce match a hand calculation done here in Python."""

import re
import sys
import openpyxl

PATH = "/tmp/claude-0/-home-user/de029611-12c7-500a-9a8b-f4dd5e8389f7/scratchpad/build/Portfolio_Tracker.xlsx"
wb = openpyxl.load_workbook(PATH)
fails, warns = [], []

SHEETS = set(wb.sheetnames)
CELL = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!\$?([A-Z]{1,3})\$?(\d+)")

# ---------------------------------------------------------------- 1. references
ref_count = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            for m in CELL.finditer(c.value):
                sheet = m.group(1) or m.group(2)
                if sheet.upper() in ("IF", "AND", "OR", "SUM"):
                    continue
                ref_count += 1
                if sheet not in SHEETS:
                    fails.append(f"{ws.title}!{c.coordinate} -> unknown sheet '{sheet}'")
print(f"[1] cross-sheet references resolved: {ref_count}")

# ---------------------------------------------------- 2. banned / risky functions
# \b so SUMIFS/COUNTIFS do not read as the post-2007 IFS()
BANNED = re.compile(
    r"\b(XLOOKUP|XMATCH|SORT|FILTER|UNIQUE|SEQUENCE|TEXTJOIN|IFS|SWITCH"
    r"|MAXIFS|MINIFS|STOCKHISTORY|LET|LAMBDA)\s*\(")
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                for m in BANNED.finditer(c.value.upper()):
                    fails.append(f"{ws.title}!{c.coordinate} uses {m.group(1)}()")
print("[2] no spilling / post-2007 unprefixed functions")

# --------------------------------------------- 3. label-and-value row alignment
def labelled(sheet, row, expect):
    got = wb[sheet][f"A{row}"].value
    if got != expect:
        fails.append(f"{sheet}!A{row} is '{got}', expected '{expect}'")

for t in ("AMZN", "MSFT", "UBER"):
    labelled(t, 8, "Shares held")
    labelled(t, 9, "Total invested (GBP)")
    labelled(t, 14, "Market value (GBP)")
    labelled(t, 15, "Unrealised P/L (GBP)")
labelled("Dashboard", 6, "Current market value (GBP)")
print("[3] holding-tab row labels line up with the formulas that target them")

# weight must divide by Dashboard market value, not any other summary row
for t in ("AMZN", "MSFT", "UBER"):
    f = wb[t]["B19"].value
    if "Dashboard!$B$6" not in f:
        fails.append(f"{t}!B19 weight divides by the wrong Dashboard row: {f}")
print("[4] % of portfolio divides by Dashboard total market value")

# --------------------------------------------------- 5. recompute the arithmetic
GBPUSD = 1.34499
EXPECT = {   # ticker: (invested, value, shares) straight from the pie export
    "AMZN": (12204.39, 13618.02, 66.0103),
    "MSFT": (5800.47, 7250.23, 20.4311),
    "UBER": (5608.26, 5831.61, 109.6538),
}

tx = wb["Transactions"]
TFIRST, TLAST = 4, 63
book = {}
for r in range(TFIRST, TLAST + 1):
    tkr = tx[f"B{r}"].value
    if not tkr:
        continue
    shares = tx[f"D{r}"].value
    px_native = tx[f"E{r}"].value
    fx = tx[f"F{r}"].value
    fee = tx[f"I{r}"].value or 0
    px_gbp = px_native / fx                      # col G
    gross = shares * px_gbp                      # col H
    net = gross + fee                            # col J, BUY
    s, cost = book.get(tkr, (0.0, 0.0))
    book[tkr] = (s + shares, cost + net)

print("\n[5] recomputing what Excel will show, from the transaction log:\n")
print(f"    {'':6} {'shares':>12} {'invested':>12} {'avg cost':>10} "
      f"{'price GBP':>10} {'mkt value':>12} {'P/L':>11} {'P/L %':>8}")
tot_inv = tot_mv = 0.0
for t, (inv_x, val_x, sh_x) in EXPECT.items():
    shares, invested = book[t]
    px_usd = wb["Prices"][f"D{ {'AMZN':7,'MSFT':8,'UBER':9}[t] }"].value
    px_gbp = px_usd / GBPUSD
    mv = shares * px_gbp
    pl = mv - invested
    tot_inv += invested
    tot_mv += mv
    print(f"    {t:6} {shares:12.4f} {invested:12.2f} {invested/shares:10.4f} "
          f"{px_gbp:10.4f} {mv:12.2f} {pl:11.2f} {pl/invested:8.2%}")

    if abs(shares - sh_x) > 1e-6:
        fails.append(f"{t} shares {shares} != export {sh_x}")
    if abs(invested - inv_x) > 0.01:
        fails.append(f"{t} invested {invested:.2f} != export {inv_x}")
    if abs(mv - val_x) > 0.01:
        fails.append(f"{t} market value {mv:.2f} != export {val_x}")

print(f"    {'TOTAL':6} {'':12} {tot_inv:12.2f} {'':10} {'':10} "
      f"{tot_mv:12.2f} {tot_mv-tot_inv:11.2f} {(tot_mv-tot_inv)/tot_inv:8.2%}")

exp_inv = sum(v[0] for v in EXPECT.values())
exp_val = sum(v[1] for v in EXPECT.values())
if abs(tot_inv - exp_inv) > 0.02 or abs(tot_mv - exp_val) > 0.02:
    fails.append(f"totals drift: {tot_inv:.2f}/{tot_mv:.2f} vs {exp_inv:.2f}/{exp_val:.2f}")

# ------------------------------------------------- 6. sell-signal logic, dry run
print("\n[6] sell-signal dry run at today's prices (defaults 50/25/40/20/25):\n")
tp, t1, t2, sl = 0.50, 0.25, 0.40, 0.20
for t in EXPECT:
    shares, invested = book[t]
    avg = invested / shares
    px_usd = wb["Prices"][f"D{ {'AMZN':7,'MSFT':8,'UBER':9}[t] }"].value
    px = px_usd / GBPUSD
    if px <= avg * (1 - sl):      sig = "SELL - STOP LOSS"
    elif px >= avg * (1 + tp):    sig = "SELL - TARGET HIT"
    elif px >= avg * (1 + t2):    sig = "TRIM 2"
    elif px >= avg * (1 + t1):    sig = "TRIM 1"
    else:                         sig = "HOLD"
    print(f"    {t:6} price {px:8.2f}  avg cost {avg:8.2f}  "
          f"stop {avg*(1-sl):7.2f}  trim1 {avg*(1+t1):7.2f}  -> {sig}")

# ---------------------------------------------------------- 7. blank-by-design
pr = wb["Prices"]
for r in (7, 8, 9):
    for col, name in (("K", "EPS TTM"), ("L", "EPS fwd"), ("O", "analyst target"),
                      ("H", "52w high"), ("I", "52w low")):
        if pr[f"{col}{r}"].value is not None:
            warns.append(f"Prices!{col}{r} ({name}) is pre-filled — should be blank")
print("\n[7] EPS / 52-week / analyst cells left blank for the user to fill")

# ------------------------------------------------------------------- verdict
print()
for w in warns:
    print("WARN:", w)
if fails:
    print(f"\nFAILED ({len(fails)}):")
    for f in fails:
        print("  -", f)
    sys.exit(1)
print("ALL CHECKS PASSED")
